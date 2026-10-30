import dataclasses
import io
import json

import httpx
from openpyxl import load_workbook

from app.bolna import BolnaError, build_agent_payload, call_user_data
from app.export import build_results_workbook
from app.payments import RazorpayPayments


def test_export_neutralises_formulas():
    data = build_results_workbook({"business_name": "=cmd()"},
                                  [{"name": "=HYPERLINK(\"evil\")", "phone": "+919876543210", "summary": "+cmd"}])
    wb = load_workbook(io.BytesIO(data))
    ws = wb["Results"]
    assert ws["B2"].value.startswith("'=")
    assert ws["C2"].value == "+919876543210"
    assert ws["G2"].value == "'+cmd"
    summary = {r[0]: r[1] for r in wb["Summary"].iter_rows(values_only=True) if r[0]}
    assert summary["Business name"] == "'=cmd()"


def test_agent_payload_from_customer_details(settings):
    campaign = {"id": "abcdefgh123", "business_name": "Acme", "business_description": "Solar since 2015",
                "product_details": "3kW system", "sales_instructions": "Book a visit",
                "welcome_message": "Hi {customer_name}, {business_name} here", "language": "Hindi"}
    voice = {"id": "v2", "name": "Kajal", "provider": "polly", "provider_config": {"voice": "Kajal", "engine": "neural"}}
    payload = build_agent_payload(settings.bolna_agent_template, campaign, "https://x/hook", voice)
    cfg = payload["agent_config"]
    llm = cfg["tasks"][0]["tools_config"]["llm_agent"]["llm_config"]
    assert "extraction_details" not in llm  # results come from Bolna Extractions (see add_extractions)
    assert llm["summarization_details"] and llm["request_json"] is False
    synth = cfg["tasks"][0]["tools_config"]["synthesizer"]
    assert synth["provider"] == "polly" and synth["provider_config"] == {"voice": "Kajal", "engine": "neural",
                                                                        "language": "hi"}
    assert cfg["agent_welcome_message"] == "Hi {{customer_name}}, Acme here"
    assert cfg["webhook_url"] == "https://x/hook" and cfg["call_summary_enabled"] is True
    transcriber = cfg["tasks"][0]["tools_config"]["transcriber"]
    # Sarvam detects the language the customer speaks, so the call can follow them.
    assert (transcriber["provider"], transcriber["model"], transcriber["language"]) == ("sarvam", "saaras:v3", "unknown")
    prompt = payload["agent_prompts"]["task_1"]["system_prompt"]
    assert "Your name is Kajal" in prompt and "Sounding natural in Hindi" in prompt
    for text in ("Acme", "Solar since 2015", "3kW system", "Book a visit", "Start the call in Hindi",
                 "{{customer_name}}", "{{customer_details}}", "How to sound human", "AI assistant calling for Acme"):
        assert text in prompt
    campaign["welcome_message"] = ""
    assert build_agent_payload(settings.bolna_agent_template, campaign, None)["agent_config"][
        "agent_welcome_message"] == "नमस्ते {{customer_name}} जी, मैं Acme से बोल रहा हूँ। दो मिनट बात कर सकते हैं?"


def test_call_user_data():
    campaign = {"detail_fields": [{"key": "city", "label": "City"}, {"key": "plan", "label": "Plan"}]}
    data = call_user_data(campaign, {"name": "Asha", "details": {"city": "Chennai", "plan": ""}})
    assert data == {"customer_name": "Asha", "customer_details": "City: Chennai", "city": "Chennai", "plan": ""}
    assert call_user_data({}, {"name": "Ravi"}) == {"customer_name": "Ravi", "customer_details": "none"}


def test_razorpay_payment_link(settings):
    seen = []

    def handler(request: httpx.Request) -> httpx.Response:
        seen.append(request)
        if request.method == "POST":
            body = json.loads(request.content)
            assert body["amount"] == 50000 and body["currency"] == "INR"
            assert body["reference_id"] == "camp1" and body["callback_method"] == "get"
            assert body["callback_url"] == "https://x/dashboard?topup=camp1"
            return httpx.Response(200, json={"id": "plink_1", "short_url": "https://rzp.io/i/x"})
        status = "paid" if request.url.path.endswith("plink_paid") else "created"
        return httpx.Response(200, json={"id": "x", "status": status, "reference_id": "camp1"})

    rp = RazorpayPayments(dataclasses.replace(settings, razorpay_key_id="rzp_test_x", razorpay_key_secret="s"))
    rp._http = httpx.Client(base_url="https://api.razorpay.com/v1", transport=httpx.MockTransport(handler),
                            auth=("rzp_test_x", "s"))
    assert rp.create_checkout("camp1", 50000, "Top-up", "https://x/dashboard?topup=camp1") == ("https://rzp.io/i/x", "plink_1")
    assert seen[0].headers["authorization"].startswith("Basic ")
    assert rp.is_paid("plink_paid", "camp1") is True
    assert rp.is_paid("plink_paid", "other-campaign") is False
    assert rp.is_paid("plink_open", "camp1") is False


def test_voices_file_is_valid(settings):
    from app.bolna import load_voices
    voices = load_voices(settings.voices_path)
    assert [v["name"] for v in voices] == ["Arjun", "Kavya"]
    assert len({v["id"] for v in voices}) == len(voices)
    for v in voices:
        assert v["name"] and v["provider"] and v["provider_config"]


def test_supabase_token_checks(settings):
    import pytest
    from app.supabase_auth import AuthError, SupabaseAuth

    def handler(request: httpx.Request) -> httpx.Response:
        assert request.url.path == "/auth/v1/user" and request.headers["apikey"] == "anon"
        token = request.headers["authorization"].removeprefix("Bearer ")
        user = {"id": "u-123", "email": "asha@gmail.com", "email_confirmed_at": "2026-10-01T10:00:00Z",
                "user_metadata": {"full_name": "Asha K"}}
        if token == "good":
            return httpx.Response(200, json=user)
        if token == "unconfirmed":
            return httpx.Response(200, json={**user, "email_confirmed_at": None})
        return httpx.Response(401, json={"msg": "invalid JWT"})

    auth = SupabaseAuth(settings)
    auth._http = httpx.Client(base_url="https://test.supabase.co", headers={"apikey": "anon"},
                              transport=httpx.MockTransport(handler))
    assert auth.verify("good") == {"sub": "u-123", "email": "asha@gmail.com", "name": "Asha K"}
    with pytest.raises(AuthError, match="confirm your email"):
        auth.verify("unconfirmed")
    with pytest.raises(AuthError, match="expired"):
        auth.verify("expired")


def test_accounts_db_from_previous_version_upgrades(tmp_path):
    import sqlite3
    from app.accounts import Accounts
    path = tmp_path / "old.db"
    db = sqlite3.connect(path)  # schema used before Google sign-in
    db.executescript("""CREATE TABLE users (id TEXT PRIMARY KEY, email TEXT NOT NULL UNIQUE, name TEXT NOT NULL,
                        password_hash TEXT NOT NULL, balance_paise INTEGER NOT NULL DEFAULT 0, created_at TEXT NOT NULL);
                        INSERT INTO users VALUES ('u1', 'old@gmail.com', 'Old', 'scrypt$x$y', 2500, '2026-01-01');""")
    db.commit()
    db.close()
    from app.db import Database
    accounts = Accounts(Database(str(path)))
    linked = accounts.user_for_identity("a-1", "OLD@gmail.com", "Old User")
    assert linked["id"] == "u1" and linked["balance_paise"] == 2500  # existing wallet kept
    new = accounts.user_for_identity("a-2", "new@gmail.com", "New")
    assert new["email"] == "new@gmail.com" and new["balance_paise"] == 0
    assert accounts.user_for_identity("a-2", "new@gmail.com", "New")["id"] == new["id"]


def test_each_language_gets_a_voice_and_transcriber_that_speak_it(settings):
    from app.bolna import LANGUAGES, load_voices
    voices = {v["id"]: v for v in load_voices(settings.voices_path)}
    base = {"id": "abcdefgh123", "business_name": "Acme", "product_details": "p", "sales_instructions": "s"}

    def tools(language, voice_id):
        payload = build_agent_payload(settings.bolna_agent_template, {**base, "language": language},
                                      None, voices[voice_id])
        t = payload["agent_config"]["tasks"][0]["tools_config"]
        return t["synthesizer"], t["transcriber"], payload["agent_prompts"]["task_1"]["system_prompt"]

    # One male and one female Sarvam voice, each speaking every call language; the customer's choice decides.
    for voice_id, speaker in (("arjun", "shubh"), ("kavya", "simran")):
        for language, code in LANGUAGES.items():
            synth, transcriber, prompt = tools(language, voice_id)
            assert synth["provider"] == "sarvam" and synth["provider_config"] == {
                "voice": speaker.title(), "voice_id": speaker, "model": "bulbul:v3", "speed": 1,
                "language": code}
            # Language switching: one Sarvam transcriber that detects the customer's language; the voice starts in
            # the campaign language (Bolna moves it to the detected one after each reply). No per-language pools.
            assert (transcriber["provider"], transcriber["model"], transcriber["language"]) == (
                "sarvam", "saaras:v3", "unknown")
            assert "multilingual" not in transcriber and "multilingual" not in synth
            assert "high_vad_sensitivity" not in transcriber
            assert f"Start the call in {language}" in prompt and "your voice follows automatically" in prompt
            assert "switch_language" not in prompt
            assert "never in English letters" in prompt and "Tamil script" in prompt
            # Switching off: fixed language, Deepgram for English and Sarvam for Indian languages.
            fixed = build_agent_payload(settings.bolna_agent_template, {**base, "language": language}, None,
                                        voices[voice_id], auto_language=False)
            fixed_tools = fixed["agent_config"]["tasks"][0]["tools_config"]
            expected = ("deepgram", "en") if language == "English" else ("sarvam", code)
            assert (fixed_tools["transcriber"]["provider"], fixed_tools["transcriber"]["language"]) == expected
            assert fixed_tools["synthesizer"]["provider_config"]["language"] == code
    # The colloquial guide covers every language the call may switch to.
    tamil_prompt = tools("Tamil", "kavya")[2]
    assert "pechu Tamil" in tamil_prompt and "Hinglish" in tamil_prompt and "spoken Telugu" in tamil_prompt


def test_agent_retried_with_region_language_codes_if_bolna_asks(settings):
    from app.bolna import BolnaClient, load_voices
    sent = []

    def handler(request: httpx.Request) -> httpx.Response:
        if request.url.path != "/v2/agent":  # extraction setup after the agent is created
            return httpx.Response(201, json={"id": "x"})
        tools = json.loads(request.content)["agent_config"]["tasks"][0]["tools_config"]
        sent.append((tools["transcriber"]["language"], tools["synthesizer"]["provider_config"]["language"]))
        if tools["synthesizer"]["provider_config"]["language"] == "ta":
            return httpx.Response(400, json={"message": "Provided language: ta is not available for the model: bulbul:v3"})
        return httpx.Response(200, json={"agent_id": "agent-9"})

    client = BolnaClient(dataclasses.replace(settings, bolna_api_key="k"))
    client._http = httpx.Client(base_url="https://api.bolna.ai", transport=httpx.MockTransport(handler))
    campaign = {"id": "abcdefgh123", "business_name": "Acme", "product_details": "p", "sales_instructions": "s",
                "language": "Tamil"}
    assert client.create_agent(campaign, None, load_voices(settings.voices_path)[1]) == "agent-9"
    assert sent == [("unknown", "ta"), ("unknown", "ta-IN")]

    # If Bolna refuses language auto-detection, a fixed-language agent is created.
    shapes = []

    def no_switching(request: httpx.Request) -> httpx.Response:
        if request.url.path != "/v2/agent":  # extraction setup after the agent is created
            return httpx.Response(201, json={"id": "x"})
        language = json.loads(request.content)["agent_config"]["tasks"][0]["tools_config"]["transcriber"]["language"]
        shapes.append(language)
        if language == "unknown":
            return httpx.Response(400, json={"message": "Provided language: unknown is not available"})
        return httpx.Response(200, json={"agent_id": "agent-10"})

    client._http = httpx.Client(base_url="https://api.bolna.ai", transport=httpx.MockTransport(no_switching))
    assert client.create_agent(campaign, None, load_voices(settings.voices_path)[1]) == "agent-10"
    assert shapes == ["unknown", "unknown", "ta"]

    # Errors that aren't about the agent's settings are reported straight away.
    sent.clear()
    client._http = httpx.Client(base_url="https://api.bolna.ai", transport=httpx.MockTransport(
        lambda r: (sent.append(1), httpx.Response(401, json={"message": "Invalid API key"}))[1]))
    import pytest
    with pytest.raises(BolnaError, match="Invalid API key"):
        client.create_agent(campaign, None, load_voices(settings.voices_path)[1])
    assert sent == [1]


def test_falls_back_to_a_smaller_ai_model_if_bolna_lacks_it(settings):
    from app.bolna import BolnaClient, load_voices
    models = []

    def handler(request: httpx.Request) -> httpx.Response:
        if request.url.path != "/v2/agent":  # extraction setup after the agent is created
            return httpx.Response(201, json={"id": "x"})
        model = json.loads(request.content)["agent_config"]["tasks"][0]["tools_config"]["llm_agent"]["llm_config"]["model"]
        models.append(model)
        if model == "gpt-4.1":
            return httpx.Response(400, json={"message": "Model gpt-4.1 is not available on your plan"})
        return httpx.Response(200, json={"agent_id": "agent-11"})

    client = BolnaClient(dataclasses.replace(settings, bolna_api_key="k"))
    client._http = httpx.Client(base_url="https://api.bolna.ai", transport=httpx.MockTransport(handler))
    campaign = {"id": "abcdefgh123", "business_name": "Acme", "product_details": "p", "sales_instructions": "s",
                "language": "Tamil"}
    assert client.create_agent(campaign, None, load_voices(settings.voices_path)[0]) == "agent-11"
    assert models == ["gpt-4.1"] * 4 + ["gpt-4o-mini"]


def test_sales_script_shows_colloquial_examples_in_native_script(settings):
    from app.bolna import COLLOQUIAL_EXAMPLES, LANGUAGES, build_system_prompt
    assert set(COLLOQUIAL_EXAMPLES) == set(LANGUAGES)
    campaign = {"business_name": "Acme", "product_details": "p", "sales_instructions": "s", "language": "Tamil"}
    prompt = build_system_prompt(campaign, "Kavya", switch_languages=True)
    assert "never like a newsreader" in prompt
    assert 'Not: "உங்களிடம் சிறிது நேரம் இருக்கிறதா?"  ->  Say: "ஒரு ரெண்டு நிமிஷம் பேசலாமா?"' in prompt
    for examples in COLLOQUIAL_EXAMPLES.values():  # every language the call may switch to
        assert examples[0][1] in prompt
    template = json.load(open(settings.bolna_agent_template))
    llm = template["agent_config"]["tasks"][0]["tools_config"]["llm_agent"]["llm_config"]
    assert llm["model"] == "gpt-4.1"


def test_opening_line_and_check_in_are_in_the_call_language(settings):
    from app.bolna import STILL_THERE, build_welcome_message, load_voices
    arjun, kavya = load_voices(settings.voices_path)
    base = {"id": "abcdefgh123", "business_name": "Sunrise Solar", "product_details": "p", "sales_instructions": "s"}
    tamil = build_welcome_message({**base, "language": "Tamil"}, kavya)
    assert tamil.startswith("வணக்கம் {{customer_name}}") and "Kavya பேசுறேன்" in tamil and "Sunrise Solar" in tamil
    assert "बोल रही हूँ" in build_welcome_message({**base, "language": "Hindi"}, kavya)   # she
    assert "बोल रहा हूँ" in build_welcome_message({**base, "language": "Hindi"}, arjun)   # he
    assert "बोलतेय" in build_welcome_message({**base, "language": "Marathi"}, kavya)
    assert build_welcome_message({**base, "language": "English"}, arjun) == (
        "Hi {{customer_name}}, this is Arjun from Sunrise Solar. Got two minutes?")
    # No agent name: no placeholder or repeated business name.
    hindi = build_welcome_message({**base, "language": "Hindi"})
    assert "{agent}" not in hindi and hindi.count("Sunrise Solar") == 1
    # The customer's own opening line is used as written.
    assert build_welcome_message({**base, "language": "Tamil", "welcome_message": "Hi {name}"}, kavya) == "Hi {{name}}"

    payload = build_agent_payload(settings.bolna_agent_template, {**base, "language": "Tamil"}, None, kavya)
    task = payload["agent_config"]["tasks"][0]
    assert task["task_config"]["check_user_online_message"] == STILL_THERE["ta"]  # in the campaign language
    assert task["task_config"]["voicemail"] is False and task["task_config"]["backchanneling"] is False
    assert task["task_config"]["hangup_after_silence"] >= 20 and task["task_config"]["call_terminate"] >= 600
    assert task["tools_config"]["transcriber"]["endpointing"] >= 500
    single = build_agent_payload(settings.bolna_agent_template, {**base, "language": "Tamil"}, None, kavya,
                                 auto_language=False)
    assert single["agent_config"]["tasks"][0]["task_config"]["check_user_online_message"] == "ஹலோ, கேக்குதா?"


def test_result_questions_attached_to_each_new_agent(settings):
    from app.bolna import EXTRACTION_CATEGORY, BolnaClient, load_voices
    calls = []

    def handler(request: httpx.Request) -> httpx.Response:
        body = json.loads(request.content) if request.content else None
        calls.append((request.method, request.url.path, body))
        if request.url.path == "/v2/agent":
            return httpx.Response(201, json={"agent_id": "agent-7"})
        if request.url.path == "/agent/agent-7/extraction-categories":
            return httpx.Response(201, json={"id": "cat-1", "name": body["name"]})
        return httpx.Response(201, json={"id": "disp"})

    client = BolnaClient(dataclasses.replace(settings, bolna_api_key="k"))
    client._http = httpx.Client(base_url="https://api.bolna.ai", transport=httpx.MockTransport(handler))
    campaign = {"id": "abcdefgh123", "business_name": "Acme", "product_details": "p", "sales_instructions": "s",
                "language": "Tamil"}
    assert client.create_agent(campaign, None, load_voices(settings.voices_path)[1]) == "agent-7"
    assert [path for _, path, _ in calls] == ["/v2/agent", "/agent/agent-7/extraction-categories"] + ["/dispositions/"] * 3
    assert calls[1][2] == {"name": EXTRACTION_CATEGORY, "model": "gpt-4.1-mini"}
    dispositions = {body["name"]: body for _, path, body in calls if path == "/dispositions/"}
    assert set(dispositions) == {"outcome", "callback_time", "key_points"}
    for body in dispositions.values():
        assert body["agent_id"] == "agent-7" and body["category_id"] == "cat-1" and body["question"]
        assert body["is_subjective"] or body["is_objective"]
    assert dispositions["outcome"]["objective_options"] == ["interested", "not_interested", "wrong_number"]
    # The retired agent-extraction settings are no longer sent.
    agent = calls[0][2]["agent_config"]
    assert [t["task_type"] for t in agent["tasks"]] == ["conversation"]
    assert "extraction_details" not in agent["tasks"][0]["tools_config"]["llm_agent"]["llm_config"]


def test_campaign_still_starts_if_bolna_refuses_the_result_questions(settings, caplog):
    import logging
    from app.bolna import BolnaClient, load_voices

    def handler(request: httpx.Request) -> httpx.Response:
        if request.url.path == "/v2/agent":
            return httpx.Response(201, json={"agent_id": "agent-8"})
        return httpx.Response(403, json={"message": "Extractions are not enabled on your plan"})

    client = BolnaClient(dataclasses.replace(settings, bolna_api_key="k"))
    client._http = httpx.Client(base_url="https://api.bolna.ai", transport=httpx.MockTransport(handler))
    campaign = {"id": "abcdefgh123", "business_name": "Acme", "product_details": "p", "sales_instructions": "s",
                "language": "English"}
    with caplog.at_level(logging.WARNING, logger="app.bolna"):
        assert client.create_agent(campaign, None, load_voices(settings.voices_path)[0]) == "agent-8"
    assert any("extractions not set up" in r.getMessage() and "not enabled" in r.getMessage() for r in caplog.records)


def test_extraction_results_fill_the_columns():
    from app.worker import parse_execution
    fields = parse_execution({"status": "completed", "transcript": "assistant: hi\nuser: sari, Saturday vaanga",
                              "extracted_data": {"outcome": "interested", "callback_time": "Saturday 10 am",
                                                 "key_points": "Asked about EMI"}})
    assert (fields["outcome"], fields["callback_time"], fields["key_points"]) == (
        "interested", "Saturday 10 am", "Asked about EMI")


def test_bolna_extractions_format_fills_the_columns():
    from app.worker import parse_execution
    fields = parse_execution({"status": "completed", "transcript": "assistant: hi\nuser: Saturday 10 vaanga",
                              "extracted_data": {"Call results": {
                                  "outcome": {"subjective": "She wants a visit", "objective": "interested",
                                              "confidence": 0.9},
                                  "callback_time": {"subjective": "Saturday 10 am", "objective": None},
                                  "key_points": {"subjective": "Asked about EMI", "objective": None}}}})
    assert (fields["outcome"], fields["callback_time"], fields["key_points"]) == (
        "interested", "Saturday 10 am", "Asked about EMI")
    none = parse_execution({"status": "completed", "transcript": "user: no thanks",
                            "extracted_data": {"Call results": {"outcome": {"objective": "not_interested"},
                                                                "callback_time": {"subjective": "none"}}}})
    assert none["outcome"] == "not_interested" and none["callback_time"] is None
