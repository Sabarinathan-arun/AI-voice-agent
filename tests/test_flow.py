import io
import json
from datetime import timedelta

from openpyxl import load_workbook

from app.store import utcnow
from app.worker import call_cost_paise, parse_execution
from tests.conftest import signup, top_up

SETUP = {"business_name": "Acme Solar", "product_details": "Rooftop solar, 3kW from 1.5L",
         "sales_instructions": "Book a free site visit.", "language": "Tamil",
         "welcome_message": "Vanakkam {customer_name}", "voice": "kavya"}
CSV = b"name,phone,City,Last Purchase\nAsha,9876543210,Chennai,Inverter\nRavi,9876543211,Madurai,\nMeena,9876543212,Salem,Panels\n"


def ready_campaign(client, csv=CSV):
    cid = client.post("/api/campaigns").json()["campaign_id"]
    assert client.put(f"/api/campaigns/{cid}/setup", json=SETUP).status_code == 200
    assert client.post(f"/api/campaigns/{cid}/leads", files={"file": ("leads.csv", csv, "text/csv")}).status_code == 200
    return cid


def test_sign_in_and_logout(client):
    assert client.get("/api/me").status_code == 401
    me = signup(client)
    assert me["email"] == "owner@example.com" and me["name"] == "Owner" and me["balance_paise"] == 0
    assert client.post("/api/auth/session", json={"access_token": "forged-token-123"}).status_code == 401
    res = client.post("/api/auth/session", json={"access_token": "unverified:new@example.com"})
    assert res.status_code == 401 and "confirm your email" in res.json()["detail"]
    client.post("/api/auth/logout")
    assert client.get("/api/me").status_code == 401
    assert signup(client)["id"] == me["id"]  # same account -> same wallet


def test_topup_credited_once(client, payments):
    signup(client)
    assert client.post("/api/wallet/topups", json={"amount_rupees": 50}).status_code == 422  # below ₹100 minimum
    res = client.post("/api/wallet/topups", json={"amount_rupees": 500})
    assert res.status_code == 201 and res.json()["checkout_url"].startswith("https://rzp.io/")
    assert payments.links[0]["amount"] == 50000 and "/dashboard?topup=" in payments.links[0]["callback"]
    assert client.post("/api/wallet/topups/verify").json()["balance_paise"] == 0  # not paid yet
    payments.pay_all()
    assert client.post("/api/wallet/topups/verify").json()["balance_paise"] == 50000
    assert client.post("/api/wallet/topups/verify").json()["balance_paise"] == 50000  # never twice
    history = client.get("/api/wallet").json()["history"]
    assert history[0]["type"] == "topup" and history[0]["amount_paise"] == 50000


def test_full_workflow_with_billing(client, bolna, payments):
    signup(client)
    cid = ready_campaign(client)
    res = client.post(f"/api/campaigns/{cid}/start")
    assert res.status_code == 402 and "₹20" in res.json()["detail"]  # below the low-balance limit
    top_up(client, payments, 500)

    res = client.post(f"/api/campaigns/{cid}/start")
    assert res.status_code == 200 and res.json()["campaign"]["status"] == "running"
    assert "bolna_agent_id" not in res.json()["campaign"] and "user_id" not in res.json()["campaign"]
    agent = bolna.agents[0]
    assert agent["language"] == "Tamil" and agent["_voice"]["id"] == "kavya"  # the agent the customer chose
    assert [v["name"] for v in client.get("/api/voices").json()] == ["Arjun", "Kavya"]

    client.worker.tick()
    assert len(bolna.calls) == 2  # concurrency limit
    agent_id, phone, data = bolna.calls[0]
    assert data == {"customer_name": "Asha", "customer_details": "City: Chennai; Last Purchase: Inverter",
                    "city": "Chennai", "last_purchase": "Inverter"}

    bolna.executions["exec-0"] = {
        "id": "exec-0", "status": "completed", "transcript": "assistant: hi\nuser: yes please",
        "summary": "Asha booked a site visit.",
        "extracted_data": {"General": {"Outcome": {"value": "Interested"}, "Callback time": {"value": "Saturday 10 am"},
                                       "Key points": {"value": "Asked about subsidy"}}},
        "conversation_duration": 90, "telephony_data": {"recording_url": "https://rec/0"},
    }
    bolna.executions["exec-1"] = {"id": "exec-1", "status": "no-answer"}
    client.worker.tick()
    # webhook for the same call must not bill twice
    client.post("/api/webhooks/bolna", json={"id": "exec-0"})
    bolna.executions["exec-2"] = {"id": "exec-2", "status": "completed", "transcript": "assistant: hello",
                                  "conversation_duration": 7}
    client.worker.tick()

    data = client.get(f"/api/campaigns/{cid}").json()
    assert data["campaign"]["status"] == "completed"
    assert data["stats"]["answered_calls"] == 2 and data["stats"]["seconds"] == 97
    assert data["stats"]["cost_paise"] == 1500 + 117  # 90 s = ₹15.00, 7 s = ₹1.17 (rounded up)
    assert data["balance_paise"] == 50000 - 1617
    by_name = {l["name"]: l for l in data["leads"]}
    assert by_name["Ravi"]["cost_paise"] == 0  # unanswered calls are free
    assert by_name["Asha"]["outcome"] == "interested" and by_name["Asha"]["callback_time"] == "Saturday 10 am"
    assert by_name["Meena"]["outcome"] == "no_answer"  # picked up but never spoke
    assert data["stats"]["booked"] == 1
    assert bolna.deleted == ["agent-1"]

    history = client.get("/api/wallet").json()["history"]
    assert any(h["type"] == "campaign" and h["amount_paise"] == -1617 and h["calls"] == 2 for h in history)

    res = client.get(f"/api/campaigns/{cid}/export")
    assert res.status_code == 200 and "Acme-Solar" in res.headers["content-disposition"]
    wb = load_workbook(io.BytesIO(res.content))
    assert wb.sheetnames == ["Summary", "Results", "Transcripts"]
    summary = {r[0]: r[1] for r in wb["Summary"].iter_rows(values_only=True) if r[0]}
    assert summary["Business name"] == "Acme Solar" and summary["AI voice"] == "Kavya"
    assert summary["Call language"] == "Tamil" and summary["Total leads"] == 3
    assert summary["Total cost (₹)"] == 16.17
    rows = list(wb["Results"].iter_rows(values_only=True))
    assert rows[0] == ("#", "Name", "Phone", "City", "Last Purchase", "Call status", "Outcome", "Call booked for",
                       "Summary", "Key points", "Talk time", "Cost (₹)", "Called at (IST)")
    assert rows[1][1:8] == ("Asha", "+919876543210", "Chennai", "Inverter", "Answered", "Interested", "Saturday 10 am")
    assert rows[1][9] == "Asked about subsidy" and rows[1][10] == "1:30" and rows[1][11] == 15.0
    assert summary["Calls booked (time agreed)"] == 1


def test_low_balance_pauses_and_resumes(client, bolna, payments, settings):
    signup(client)
    top_up(client, payments, 100)
    cid = ready_campaign(client)
    client.post(f"/api/campaigns/{cid}/start")
    client.worker.tick()
    # a long call drains the wallet below the ₹20 low-balance limit
    bolna.executions["exec-0"] = {"id": "exec-0", "status": "completed", "conversation_duration": 492}
    bolna.executions["exec-1"] = {"id": "exec-1", "status": "busy"}
    client.worker.tick()
    data = client.get(f"/api/campaigns/{cid}").json()
    assert data["balance_paise"] == 10000 - 8200  # ₹18 left, below ₹20
    assert data["campaign"]["low_balance"] is True and len(bolna.calls) == 2  # Meena not called yet
    top_up(client, payments, 200)
    client.worker.tick()
    assert len(bolna.calls) == 3
    assert client.get(f"/api/campaigns/{cid}").json()["campaign"]["low_balance"] is False


def test_campaigns_are_private_and_autofilled(client, payments):
    signup(client)
    cid = ready_campaign(client)
    second = client.post("/api/campaigns").json()["campaign_id"]
    prefilled = client.get(f"/api/campaigns/{second}").json()["campaign"]
    assert prefilled["business_name"] == "Acme Solar" and prefilled["language"] == "Tamil"
    assert prefilled["voice"] == "kavya" and prefilled["prefilled"] is True
    assert [c["id"] for c in client.get("/api/campaigns").json()] == [second, cid]

    client.post("/api/auth/logout")
    signup(client, email="other@example.com")
    assert client.get(f"/api/campaigns/{cid}").status_code == 404
    assert client.get(f"/api/campaigns/{cid}/export").status_code == 404
    assert client.get("/api/campaigns").json() == []


def test_stop_skips_pending(client, bolna, payments):
    signup(client)
    top_up(client, payments)
    cid = ready_campaign(client)
    client.post(f"/api/campaigns/{cid}/start")
    client.worker.tick()
    assert client.post(f"/api/campaigns/{cid}/stop").json()["campaign"]["status"] == "stopped"
    for eid in ("exec-0", "exec-1"):
        bolna.executions[eid] = {"id": eid, "status": "busy"}
    client.worker.tick()
    leads = client.get(f"/api/campaigns/{cid}").json()["leads"]
    assert [l["status"] for l in leads] == ["busy", "busy", "skipped"]
    assert bolna.deleted == ["agent-1"]


def test_unknown_language_rejected(client):
    signup(client)
    cid = client.post("/api/campaigns").json()["campaign_id"]
    assert client.put(f"/api/campaigns/{cid}/setup", json={**SETUP, "language": "Klingon"}).status_code == 422


def test_call_timeout(client, bolna, payments):
    signup(client)
    top_up(client, payments)
    cid = ready_campaign(client, b"name,phone\nA,9876543210\n")
    client.post(f"/api/campaigns/{cid}/start")
    client.worker.tick()
    client.worker.tick(utcnow() + timedelta(hours=2))
    data = client.get(f"/api/campaigns/{cid}").json()
    assert data["leads"][0]["status"] == "timed_out"
    assert data["campaign"]["status"] == "completed"


def test_lead_data_deleted_seven_days_after_finish(client, bolna, payments):
    signup(client)
    top_up(client, payments)
    draft = client.post("/api/campaigns").json()["campaign_id"]           # never started
    done = ready_campaign(client, b"name,phone\nA,9876543210\n")
    client.post(f"/api/campaigns/{done}/start")
    client.worker.tick()
    bolna.executions["exec-0"] = {"id": "exec-0", "status": "completed", "transcript": "user: hi",
                                  "conversation_duration": 30}
    client.worker.tick()
    assert client.get(f"/api/campaigns/{done}").json()["campaign"]["status"] == "completed"

    client.worker._last_purge = None
    client.worker.tick(utcnow() + timedelta(days=6))
    assert client.get(f"/api/campaigns/{done}").status_code == 200       # still downloadable on day 6
    assert client.get(f"/api/campaigns/{done}/export").status_code == 200
    client.worker._last_purge = None
    client.worker.tick(utcnow() + timedelta(days=8))
    assert client.get(f"/api/campaigns/{done}").status_code == 404       # leads and results gone
    assert client.get(f"/api/campaigns/{draft}").status_code == 404
    assert client.store.get_leads(done) == []
    history = client.get("/api/wallet").json()["history"]                  # cost totals are kept
    assert any(h["type"] == "campaign" and h["amount_paise"] == -500 for h in history)
    assert client.get("/api/me").status_code == 200


def test_running_campaign_survives_a_restart(settings, bolna, payments):
    from fastapi.testclient import TestClient
    from app.main import create_app
    from tests.conftest import FakeAuth

    first = create_app(settings, bolna=bolna, payments=payments, auth=FakeAuth(), run_worker=False)
    with TestClient(first) as c:
        signup(c)
        top_up(c, payments)
        cid = ready_campaign(c)
        c.post(f"/api/campaigns/{cid}/start")
        first.state.worker.tick()                     # two calls in progress, one lead waiting
        cookies = dict(c.cookies)
    assert len(bolna.calls) == 2

    # "crash": a brand-new app on the same database carries on with the campaign
    second = create_app(settings, bolna=bolna, payments=payments, auth=FakeAuth(), run_worker=False)
    with TestClient(second, cookies=cookies) as c:
        for eid in ("exec-0", "exec-1"):
            bolna.executions[eid] = {"id": eid, "status": "completed", "transcript": "user: yes",
                                     "conversation_duration": 60, "extracted_data": {"outcome": "interested"}}
        second.state.worker.tick()
        assert len(bolna.calls) == 3                  # the waiting lead is called after the restart
        bolna.executions["exec-2"] = {"id": "exec-2", "status": "no-answer"}
        second.state.worker.tick()
        data = c.get(f"/api/campaigns/{cid}").json()
        assert data["campaign"]["status"] == "completed"
        assert [l["outcome"] for l in data["leads"]] == ["interested", "interested", "no_answer"]
        assert data["balance_paise"] == 50000 - 2000  # both answered calls billed exactly once


def test_wallet_transaction_log(client, bolna, payments):
    signup(client)
    top_up(client, payments, 200)
    cid = ready_campaign(client, b"name,phone\nA,9876543210\n")
    client.post(f"/api/campaigns/{cid}/start")
    client.worker.tick()
    bolna.executions["exec-0"] = {"id": "exec-0", "status": "completed", "transcript": "user: hi",
                                  "conversation_duration": 90}
    client.worker.tick()
    user = client.get("/api/me").json()
    log = client.accounts.transactions(user["id"])
    assert [(t["kind"], t["amount_paise"], t["balance_after_paise"]) for t in log] == [
        ("call", -1500, 18500), ("topup", 20000, 20000)]

def test_call_cost():
    assert call_cost_paise(0, 1000) == 0
    assert call_cost_paise(60, 1000) == 1000
    assert call_cost_paise(1, 1000) == 17


def test_parse_execution_variants():
    f = parse_execution({"status": "completed", "transcript": [{"role": "user", "content": "hi"}],
                         "extracted_data": json.dumps({"outcome": {"value": "callback requested"}})})
    assert f["outcome"] == "interested" and f["transcript"] == "user: hi" and f["summary"] == "user: hi"
    assert parse_execution({"status": "scheduled"})["status"] == "in-progress"
    assert "outcome" not in parse_execution({"status": "ringing"})
    assert "outcome" not in parse_execution({"status": "call-disconnected"})


def test_parse_execution_nested_extraction():
    f = parse_execution({
        "status": "completed", "conversation_duration": 16, "transcript": "assistant: hi\nuser: hello",
        "extracted_data": {"General": {"Call Summary": {"subjective": "Customer wants a demo."},
                                       "Outcome": {"value": "Interested"}}},
        "telephony_data": {"duration": 16, "recording_url": "https://api.bolna.ai/recordings/call/x"},
    })
    assert f["summary"] == "Customer wants a demo." and f["outcome"] == "interested"
    assert f["duration_seconds"] == 16 and f["recording_url"].endswith("/x")


def test_answered_calls_are_interested_or_not():
    talk = "assistant: hello\nuser: tell me"
    base = {"status": "completed", "transcript": talk}
    assert parse_execution({**base, "extracted_data": {"outcome": "not interested"}})["outcome"] == "not_interested"
    assert parse_execution({**base, "extracted_data": {"outcome": "wrong number"}})["outcome"] == "wrong_number"
    # a booked time means interested even without an outcome
    f = parse_execution({**base, "extracted_data": {"callback_time": "tomorrow 5 pm"}})
    assert f["outcome"] == "interested" and f["callback_time"] == "tomorrow 5 pm"
    assert parse_execution({**base, "extracted_data": {"callback_time": "none"}})["callback_time"] is None
    # no extraction from Bolna: judged from the summary, never "completed"
    assert parse_execution({**base, "summary": "Customer is not interested at the moment."})["outcome"] == "not_interested"
    assert parse_execution({**base, "summary": "Customer agreed to a site visit."})["outcome"] == "interested"
    assert parse_execution({**base, "summary": "Short chat."})["outcome"] == "not_interested"
    assert parse_execution({"status": "completed", "transcript": "assistant: hello"})["outcome"] == "no_answer"


def test_campaign_starts_only_once(client, bolna, payments):
    from app.bolna import BolnaError
    signup(client)
    top_up(client, payments)
    cid = ready_campaign(client)
    store = client.app.state.store
    seen = []
    real_create = bolna.create_agent

    def create_agent(campaign, webhook_url, voice=None):
        # A second click arriving while the agent is being created must not start the campaign again.
        seen.append((store.get_campaign(cid)["status"], store.transition_campaign(cid, "setup", "starting")))
        if len(seen) == 1:
            raise BolnaError("Bolna is down")
        return real_create(campaign, webhook_url, voice)

    bolna.create_agent = create_agent
    assert client.post(f"/api/campaigns/{cid}/start").status_code == 502
    assert store.get_campaign(cid)["status"] == "setup"  # failed start can be retried
    assert client.post(f"/api/campaigns/{cid}/start").json()["campaign"]["status"] == "running"
    assert seen == [("starting", False), ("starting", False)]
    assert client.post(f"/api/campaigns/{cid}/start").status_code == 409
    assert len(bolna.agents) == 1


def test_stop_during_last_call_stays_stopped(client, bolna, payments):
    signup(client)
    top_up(client, payments)
    cid = ready_campaign(client, b"name,phone\nA,9876543210\nB,9876543211\n")
    client.post(f"/api/campaigns/{cid}/start")
    store = client.app.state.store
    campaign = store.get_campaign(cid)  # the worker's copy, read while still running
    leads = store.leads_with_status(cid, ("pending",))
    client.post(f"/api/campaigns/{cid}/stop")
    for lead in leads:  # leads picked before the stop are not called
        client.worker.dial(campaign, lead, utcnow())
    assert bolna.calls == []
    client.worker.process_campaign(campaign, utcnow())
    assert client.get(f"/api/campaigns/{cid}").json()["campaign"]["status"] == "stopped"
    assert store.get_campaign(cid)["finished_at"]


def test_call_placed_but_not_saved_is_recovered_by_webhook(client, bolna, payments):
    signup(client)
    top_up(client, payments)
    cid = ready_campaign(client, b"name,phone\nA,9876543210\n")
    client.post(f"/api/campaigns/{cid}/start")
    store = client.app.state.store
    lead = store.get_leads(cid)[0]
    # The server stopped right after Bolna accepted the call: lead is "dialing" with no execution id.
    store.claim_dial(lead["id"])
    store.update_lead(lead["id"], called_at=utcnow().isoformat())
    bolna.executions["exec-x"] = {"id": "exec-x", "status": "completed", "agent_id": "agent-1",
                                  "user_number": "+919876543210", "transcript": "user: yes, interested",
                                  "conversation_duration": 60}
    assert client.post("/api/webhooks/bolna", json={"id": "exec-x"}).status_code == 200
    data = client.get(f"/api/campaigns/{cid}").json()
    assert data["leads"][0]["status"] == "completed" and data["leads"][0]["cost_paise"] == 1000
    assert data["balance_paise"] == 50000 - 1000
    # A call from some other agent or number is ignored.
    bolna.executions["exec-y"] = {"id": "exec-y", "status": "completed", "agent_id": "other",
                                  "user_number": "+919876543210"}
    assert client.post("/api/webhooks/bolna", json={"id": "exec-y"}).status_code == 200
    assert client.get("/api/me").json()["balance_paise"] == 49000


def test_sign_in_edge_cases(client):
    accounts = client.accounts
    a = accounts.user_for_identity("auth-a", "a@example.com", "A")
    b = accounts.user_for_identity("auth-b", "b@example.com", "B")
    # B changes their Supabase email to one A already uses: sign-in still works, B keeps their address.
    assert accounts.user_for_identity("auth-b", "a@example.com", "B")["email"] == "b@example.com"
    # B changes to a free address: followed.
    assert accounts.user_for_identity("auth-b", "new@example.com", "B")["email"] == "new@example.com"
    # Same person signing in again (e.g. a second tab) gets the same account.
    assert accounts.user_for_identity("auth-a", "a@example.com", "A")["id"] == a["id"]
    assert b["id"] != a["id"]


def test_campaign_list_counts(client, bolna, payments):
    signup(client)
    top_up(client, payments)
    draft = client.post("/api/campaigns").json()["campaign_id"]
    cid = ready_campaign(client)
    client.post(f"/api/campaigns/{cid}/start")
    client.worker.tick()  # 2 calls ringing, 1 waiting

    def counts():
        return {c["id"]: (c["leads"], c["done"], c["business_name"]) for c in client.get("/api/campaigns").json()}

    assert counts() == {cid: (3, 0, "Acme Solar"), draft: (0, 0, None)}  # draft made before any setup
    client.post(f"/api/campaigns/{cid}/stop")
    assert counts()[cid] == (3, 1, "Acme Solar")  # the waiting lead is skipped
    bolna.executions["exec-0"] = {"id": "exec-0", "status": "busy"}
    client.worker.tick()
    assert counts()[cid] == (3, 2, "Acme Solar")  # one call finished, one still ringing


def test_finished_call_logged_once_without_personal_data(client, bolna, payments, caplog):
    import logging
    signup(client)
    top_up(client, payments)
    cid = ready_campaign(client, b"name,phone\nAsha,9876543210\n")
    client.post(f"/api/campaigns/{cid}/start")
    client.worker.tick()
    bolna.executions["exec-0"] = {"id": "exec-0", "status": "completed", "conversation_duration": 30,
                                  "transcript": "assistant: Vanakkam Asha\nassistant: Kekkudha?\n",
                                  "telephony_data": {"hangup_reason": "silence timeout"}}
    with caplog.at_level(logging.INFO, logger="app.worker"):
        client.worker.tick()
        client.post("/api/webhooks/bolna", json={"id": "exec-0"})  # same call reported again
    lines = [r.getMessage() for r in caplog.records if "call finished" in r.getMessage()]
    assert len(lines) == 1
    assert "agent_turns=2 customer_turns=0" in lines[0] and "hangup=silence timeout" in lines[0]
    assert "Asha" not in lines[0] and "9876543210" not in lines[0]


def test_language_switching_can_be_turned_off(settings):
    import dataclasses
    import httpx
    from app.bolna import BolnaClient, load_voices
    sent = []

    def handler(request):
        sent.append("multilingual" in json.loads(request.content)["agent_config"]["tasks"][0]["tools_config"]["transcriber"])
        return httpx.Response(200, json={"agent_id": "a"})

    for switching, expected in ((True, [True]), (False, [False])):
        sent.clear()
        client = BolnaClient(dataclasses.replace(settings, bolna_api_key="k", language_switching=switching))
        client._http = httpx.Client(base_url="https://api.bolna.ai", transport=httpx.MockTransport(handler))
        client.create_agent({"id": "abcdefgh1", "business_name": "A", "product_details": "p",
                             "sales_instructions": "s", "language": "Tamil"}, None, load_voices(settings.voices_path)[1])
        assert sent == expected
