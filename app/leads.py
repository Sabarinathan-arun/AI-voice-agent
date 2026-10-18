"""Parse uploaded lead spreadsheets (.xlsx or .csv) into name/phone records."""

import csv
import io
import re
from dataclasses import dataclass, field

from openpyxl import load_workbook

NAME_HEADERS = ("name", "customer", "lead", "contact name", "full name")
PHONE_HEADERS = ("phone", "mobile", "number", "contact", "cell", "whatsapp", "tel")


class LeadFileError(ValueError):
    pass


@dataclass
class ParseResult:
    leads: list[dict] = field(default_factory=list)
    errors: list[dict] = field(default_factory=list)


def normalize_phone(raw, default_country_code: str) -> str | None:
    """Return the number in E.164 format (+<digits>) or None if it is not a plausible phone number."""
    if raw is None:
        return None
    if isinstance(raw, float) and raw.is_integer():
        raw = int(raw)
    text = str(raw).strip()
    if not text:
        return None
    has_plus = text.startswith("+")
    digits = re.sub(r"\D", "", text)
    if not has_plus:
        if digits.startswith("00"):
            digits = digits[2:]
        elif len(digits) == 11 and digits.startswith("0"):
            digits = default_country_code + digits[1:]
        elif len(digits) == 10:
            digits = default_country_code + digits
    if not 8 <= len(digits) <= 15:
        return None
    return "+" + digits


def _rows_from_xlsx(content: bytes) -> list[list]:
    try:
        wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    except Exception as exc:
        raise LeadFileError("Could not read the Excel file. Please upload a valid .xlsx file.") from exc
    ws = wb.worksheets[0]
    rows = [list(r) for r in ws.iter_rows(values_only=True)]
    wb.close()
    return rows


def _rows_from_csv(content: bytes) -> list[list]:
    try:
        text = content.decode("utf-8-sig")
    except UnicodeDecodeError:
        text = content.decode("latin-1")
    return [row for row in csv.reader(io.StringIO(text))]


def _find_column(header: list[str], candidates: tuple[str, ...]) -> int | None:
    for i, cell in enumerate(header):
        if any(c in cell for c in candidates):
            return i
    return None


def parse_leads(filename: str, content: bytes, default_country_code: str, max_leads: int) -> ParseResult:
    name = (filename or "").lower()
    if name.endswith(".xlsx"):
        rows = _rows_from_xlsx(content)
    elif name.endswith(".csv"):
        rows = _rows_from_csv(content)
    else:
        raise LeadFileError("Unsupported file type. Please upload an .xlsx (or .csv) file.")

    rows = [r for r in rows if any(c not in (None, "") for c in r)]
    if not rows:
        raise LeadFileError("The file is empty.")

    header = [str(c or "").strip().lower() for c in rows[0]]
    phone_col = _find_column(header, PHONE_HEADERS)
    name_col = _find_column([h if i != phone_col else "" for i, h in enumerate(header)], NAME_HEADERS)
    if phone_col is not None:
        data_rows, first_row_number = rows[1:], 2
        if name_col is None:
            name_col = 0 if phone_col != 0 else 1
    else:
        # No recognisable header: assume column A = name, column B = phone.
        name_col, phone_col = 0, 1
        if normalize_phone(rows[0][1] if len(rows[0]) > 1 else None, default_country_code):
            data_rows, first_row_number = rows, 1
        else:
            data_rows, first_row_number = rows[1:], 2

    result = ParseResult()
    seen: set[str] = set()
    for offset, row in enumerate(data_rows):
        row_number = first_row_number + offset
        raw_name = row[name_col] if name_col < len(row) else None
        raw_phone = row[phone_col] if phone_col < len(row) else None
        phone = normalize_phone(raw_phone, default_country_code)
        if not phone:
            result.errors.append({"row": row_number, "reason": f"Invalid phone number: {raw_phone!r}"})
            continue
        if phone in seen:
            result.errors.append({"row": row_number, "reason": f"Duplicate phone number: {phone}"})
            continue
        seen.add(phone)
        lead_name = str(raw_name).strip() if raw_name not in (None, "") else "there"
        result.leads.append({"row_number": row_number, "name": lead_name[:200], "phone": phone})

    if not result.leads:
        raise LeadFileError("No valid leads found. The file needs a name column and a phone number column.")
    if len(result.leads) > max_leads:
        raise LeadFileError(f"Too many leads ({len(result.leads)}). The maximum per campaign is {max_leads}.")
    return result
