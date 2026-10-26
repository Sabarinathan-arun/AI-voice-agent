"""Build the downloadable campaign report (.xlsx): a Summary sheet with every campaign detail, usage and cost,
a formatted Results sheet (one row per lead) and a Transcripts sheet."""

import io
import re
from datetime import datetime, timedelta, timezone

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

IST = timezone(timedelta(hours=5, minutes=30))

BRAND = "5B4BFF"
INK = "12172B"
MUTED = "5B6378"
HEADER_FILL = PatternFill("solid", fgColor=BRAND)
SECTION_FILL = PatternFill("solid", fgColor="EEECFF")
LABEL_FILL = PatternFill("solid", fgColor="F6F6F3")
THIN = Side(style="thin", color="E5E7EF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
WRAP_TOP = Alignment(wrap_text=True, vertical="top")
WRAP_TOP_LEFT = Alignment(wrap_text=True, vertical="top", horizontal="left")

OUTCOME_COLORS = {  # fill, font
    "interested": ("E6F6F0", "0F7B5F"),
    "callback_requested": ("FFF4E0", "A35A00"),
    "not_interested": ("FDECEA", "B42318"),
    "wrong_number": ("FDECEA", "B42318"),
    "no_answer": ("F1F2F6", "5B6378"),
    "busy": ("F1F2F6", "5B6378"),
    "voicemail": ("F1F2F6", "5B6378"),
    "no_conversation": ("F1F2F6", "5B6378"),
    "failed": ("FDECEA", "B42318"),
}
NOT_REACHED = {"no_answer", "busy", "failed", "cancelled", "voicemail", "unknown"}


def _safe(value):
    """Neutralise spreadsheet formula injection from call content."""
    if isinstance(value, str):
        if re.fullmatch(r"\+\d{8,15}", value):  # E.164 phone numbers are safe as-is
            return value
        if value[:1] in ("=", "+", "-", "@", "\t", "\r"):
            return "'" + value
        return value[:32000]
    return value


def _label(value) -> str:
    """'callback_requested' / 'no-answer' -> 'Callback requested' / 'No answer'."""
    text = str(value or "").replace("_", " ").replace("-", " ").strip()
    return text[:1].upper() + text[1:]


def _ist(value) -> str:
    if not value:
        return ""
    try:
        return datetime.fromisoformat(value).astimezone(IST).strftime("%d %b %Y, %I:%M %p")
    except ValueError:
        return str(value)


def _duration(seconds) -> str:
    if not seconds:
        return ""
    seconds = int(round(float(seconds)))
    return f"{seconds // 60}:{seconds % 60:02d}"


def _rupees(paise) -> float:
    return round((paise or 0) / 100, 2)


def _write_section(ws, row: int, title: str) -> int:
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=4)
    cell = ws.cell(row=row, column=1, value=title)
    cell.font = Font(bold=True, size=12, color=BRAND)
    cell.fill = SECTION_FILL
    for col in range(1, 5):
        ws.cell(row=row, column=col).fill = SECTION_FILL
    ws.row_dimensions[row].height = 22
    return row + 1


def _write_pair(ws, row: int, label: str, value, number_format: str | None = None) -> int:
    lc = ws.cell(row=row, column=1, value=label)
    lc.font = Font(bold=True, color=INK)
    lc.fill = LABEL_FILL
    lc.alignment = WRAP_TOP
    lc.border = BORDER
    ws.merge_cells(start_row=row, start_column=2, end_row=row, end_column=4)
    vc = ws.cell(row=row, column=2, value=_safe(value) if value not in (None, "") else "—")
    vc.alignment = WRAP_TOP_LEFT
    if number_format:
        vc.number_format = number_format
    for col in range(2, 5):
        ws.cell(row=row, column=col).border = BORDER
    if isinstance(value, str) and len(value) > 70:  # give long text room (Excel does not auto-size merged cells)
        ws.row_dimensions[row].height = min(15 * (len(value) // 70 + 1 + value.count("\n")), 300)
    return row + 1


def build_results_workbook(campaign: dict, leads: list[dict], usage: dict | None = None, voice_name: str = "",
                           price_per_minute_paise: int = 0) -> bytes:
    usage = usage or {"calls": 0, "seconds": 0, "amount_paise": 0}
    detail_fields = campaign.get("detail_fields") or []
    wb = Workbook()

    # ---------------- Summary ----------------
    ws = wb.active
    ws.title = "Summary"
    ws.sheet_view.showGridLines = False
    for col, width in zip("ABCD", (30, 34, 34, 34)):
        ws.column_dimensions[col].width = width

    ws.merge_cells("A1:D1")
    ws["A1"] = "Intuvora · Campaign report"
    ws["A1"].font = Font(bold=True, size=18, color="FFFFFF")
    ws["A1"].fill = HEADER_FILL
    ws["A1"].alignment = Alignment(vertical="center", indent=1)
    ws.row_dimensions[1].height = 36
    for col in "BCD":
        ws[f"{col}1"].fill = HEADER_FILL
    ws.merge_cells("A2:D2")
    ws["A2"] = f"{campaign.get('business_name') or 'Campaign'} — generated {datetime.now(IST).strftime('%d %b %Y, %I:%M %p')} IST"
    ws["A2"].font = Font(italic=True, color=MUTED)
    ws["A2"].alignment = Alignment(indent=1)

    row = 4
    row = _write_section(ws, row, "Campaign details")
    for label, value in [
        ("Business name", campaign.get("business_name")),
        ("About the business", campaign.get("business_description")),
        ("Product / offer", campaign.get("product_details")),
        ("Sales instructions", campaign.get("sales_instructions")),
        ("Opening line", campaign.get("welcome_message") or "Default greeting"),
        ("Call language", campaign.get("language")),
        ("AI voice", voice_name),
        ("Lead file", campaign.get("lead_file")),
        ("Extra lead columns used", ", ".join(f["label"] for f in detail_fields)),
        ("Status", _label(campaign.get("status"))),
        ("Created (IST)", _ist(campaign.get("created_at"))),
        ("Calls started (IST)", _ist(campaign.get("started_at"))),
        ("Calls finished (IST)", _ist(campaign.get("finished_at"))),
    ]:
        row = _write_pair(ws, row, label, value)

    outcomes: dict[str, int] = {}
    for lead in leads:
        key = lead.get("outcome") or lead.get("status") or "pending"
        outcomes[key] = outcomes.get(key, 0) + 1
    total = len(leads)
    reached = sum(1 for l in leads if l.get("outcome") and l["outcome"] not in NOT_REACHED)
    not_reached = sum(1 for l in leads if l.get("outcome") in NOT_REACHED)

    row += 1
    row = _write_section(ws, row, "Results overview")
    row = _write_pair(ws, row, "Total leads", total)
    row = _write_pair(ws, row, "Conversations (answered)", reached)
    row = _write_pair(ws, row, "Interested", sum(1 for l in leads if l.get("outcome") == "interested"))
    row = _write_pair(ws, row, "Calls booked (time agreed)", sum(1 for l in leads if l.get("callback_time")))
    row = _write_pair(ws, row, "Not reached (no answer, busy, failed)", not_reached)
    row = _write_pair(ws, row, "Answer rate", (reached / total) if total else 0, "0%")

    row += 1
    headers = ("Outcome", "Leads", "Share of leads", "")
    for col, text in enumerate(headers[:3], start=1):
        c = ws.cell(row=row, column=col, value=text)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = HEADER_FILL
        c.border = BORDER
    row += 1
    for key, count in sorted(outcomes.items(), key=lambda kv: -kv[1]):
        fill, font = OUTCOME_COLORS.get(key, ("FFFFFF", INK))
        oc = ws.cell(row=row, column=1, value=_label(key))
        oc.fill = PatternFill("solid", fgColor=fill)
        oc.font = Font(bold=True, color=font)
        ws.cell(row=row, column=2, value=count)
        sc = ws.cell(row=row, column=3, value=(count / total) if total else 0)
        sc.number_format = "0%"
        for col in range(1, 4):
            ws.cell(row=row, column=col).border = BORDER
        row += 1

    row += 1
    row = _write_section(ws, row, "Usage & cost")
    row = _write_pair(ws, row, "Answered calls billed", usage.get("calls", 0))
    row = _write_pair(ws, row, "Total talk time (min:sec)", _duration(usage.get("seconds")) or "0:00")
    rate = _rupees(price_per_minute_paise)
    row = _write_pair(ws, row, "Rate", f"₹{rate:g} per minute, billed per second. Unanswered calls are free.")
    row = _write_pair(ws, row, "Total cost (₹)", _rupees(usage.get("amount_paise")), '"₹"#,##0.00')

    row += 1
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=4)
    note = ws.cell(row=row, column=1, value="Lead details and call results are deleted from Intuvora 7 days after the "
                                            "campaign finishes. Keep this file for your records.")
    note.font = Font(italic=True, color=MUTED, size=9)
    note.alignment = WRAP_TOP
    ws.page_setup.orientation = "portrait"
    ws.page_setup.fitToWidth = 1
    ws.sheet_properties.pageSetUpPr.fitToPage = True

    # ---------------- Results ----------------
    rs = wb.create_sheet("Results")
    columns = ([("#", 6), ("Name", 22), ("Phone", 17)]
               + [(f["label"], 18) for f in detail_fields]
               + [("Call status", 15), ("Outcome", 16), ("Call booked for", 20), ("Summary", 60), ("Key points", 40),
                  ("Talk time", 11), ("Cost (₹)", 12), ("Called at (IST)", 21)])
    for col, (title, width) in enumerate(columns, start=1):
        c = rs.cell(row=1, column=col, value=_safe(title))
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = HEADER_FILL
        c.alignment = Alignment(vertical="center", wrap_text=True)
        c.border = BORDER
        rs.column_dimensions[get_column_letter(col)].width = width
    rs.row_dimensions[1].height = 28
    outcome_col = 3 + len(detail_fields) + 2
    booked_col = outcome_col + 1
    cost_col = outcome_col + 5

    for i, lead in enumerate(leads, start=1):
        details = lead.get("details") or {}
        values = ([i, lead.get("name"), lead.get("phone")]
                  + [details.get(f["key"]) for f in detail_fields]
                  + ["Answered" if lead.get("status") == "completed" else _label(lead.get("status")),
                     _label(lead.get("outcome")), lead.get("callback_time"),
                     lead.get("summary"), lead.get("key_points"), _duration(lead.get("duration_seconds")),
                     _rupees(lead.get("cost_paise")), _ist(lead.get("called_at"))])
        r = i + 1
        for col, value in enumerate(values, start=1):
            c = rs.cell(row=r, column=col, value=_safe(value) if value not in (None, "") else None)
            c.alignment = WRAP_TOP
            c.border = BORDER
        rs.cell(row=r, column=3).number_format = "@"
        rs.cell(row=r, column=cost_col).number_format = '"₹"#,##0.00'
        fill, font = OUTCOME_COLORS.get(lead.get("outcome") or "", (None, None))
        if fill:
            oc = rs.cell(row=r, column=outcome_col)
            oc.fill = PatternFill("solid", fgColor=fill)
            oc.font = Font(bold=True, color=font)
        if lead.get("callback_time"):
            rs.cell(row=r, column=booked_col).font = Font(bold=True, color="0F7B5F")
        if i % 2 == 0:
            for col in range(1, len(columns) + 1):
                cell = rs.cell(row=r, column=col)
                if not (col == outcome_col and fill):
                    cell.fill = PatternFill("solid", fgColor="FAFAFC")
    rs.freeze_panes = "D2"
    rs.auto_filter.ref = f"A1:{get_column_letter(len(columns))}{max(len(leads) + 1, 1)}"
    rs.page_setup.orientation = "landscape"
    rs.page_setup.fitToWidth = 1
    rs.page_setup.fitToHeight = 0
    rs.sheet_properties.pageSetUpPr.fitToPage = True
    rs.print_title_rows = "1:1"

    # ---------------- Transcripts ----------------
    ts = wb.create_sheet("Transcripts")
    for col, (title, width) in enumerate([("#", 6), ("Name", 22), ("Phone", 17), ("Outcome", 20), ("Transcript", 110)],
                                         start=1):
        c = ts.cell(row=1, column=col, value=title)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = HEADER_FILL
        c.border = BORDER
        ts.column_dimensions[get_column_letter(col)].width = width
    r = 2
    for i, lead in enumerate(leads, start=1):
        if not lead.get("transcript"):
            continue
        values = [i, lead.get("name"), lead.get("phone"), _label(lead.get("outcome")), lead["transcript"]]
        for col, value in enumerate(values, start=1):
            c = ts.cell(row=r, column=col, value=_safe(value))
            c.alignment = WRAP_TOP
            c.border = BORDER
        ts.cell(row=r, column=3).number_format = "@"
        ts.row_dimensions[r].height = min(15 * (lead["transcript"].count("\n") + 2), 400)
        r += 1
    if r == 2:
        ts.cell(row=2, column=1, value="No call transcripts yet.").font = Font(italic=True, color=MUTED)
    ts.freeze_panes = "A2"
    ts.page_setup.orientation = "landscape"
    ts.page_setup.fitToWidth = 1
    ts.page_setup.fitToHeight = 0
    ts.sheet_properties.pageSetUpPr.fitToPage = True

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
